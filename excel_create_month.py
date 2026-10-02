#!/usr/bin/env python3
"""
excel_create_month.py — Python replacement for excel-create-month.ps1
Dọn dẹp workbook đầu tháng: xóa các sheet đơn hàng cũ, giữ 1 sheet mẫu → đổi tên thành "0-MAU".

Usage:
    python excel_create_month.py --file <path>
"""

import argparse
import os
import re
import sys
import zipfile
import hashlib

from openpyxl import load_workbook

# Fix UnicodeEncodeError on Windows console (cp1252 can't handle Vietnamese)
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')


MAX_SHEET_NUM = 200
SHEET_PATTERN = re.compile(r'^0*(\d+)-(.*)')


def heal_with_excel(path):
    """
    Re-save the workbook through Excel COM so it is "clean" afterwards.

    openpyxl can leave a workbook in a state where Excel wants to repair it
    on open. An interactive Excel does that silently, but the background
    automation (excel_automation.py) then fails with COM error 0x800AC472
    ("Open method of Workbooks class failed"). Opening once with
    CorruptLoad=1 (xlRepairFile) and re-saving via SaveAs heals the file.

    Best-effort: returns True on success, False otherwise (the automation
    has its own repair-mode fallback, so failure here is not fatal).
    """
    if not str(path).lower().endswith('.xlsx'):
        return False  # keep legacy .xls files untouched
    try:
        import os
        import pythoncom
        import win32com.client
    except ImportError:
        print("WARN: win32com not available; skipping Excel heal step.")
        return False

    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        wb = app.Workbooks.Open(path, CorruptLoad=1)
        tmp = path + ".heal.xlsx"
        if os.path.exists(tmp):
            os.remove(tmp)
        wb.SaveAs(tmp, FileFormat=51)  # xlOpenXMLWorkbook
        wb.Close(SaveChanges=False)
        os.replace(tmp, path)
        print("Healed workbook via Excel re-save.")
        return True
    except Exception as e:
        print(f"WARN: Excel heal skipped: {e}")
        return False
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def deduplicate_media_images(xlsx_path, verbose=True):
    """Khử trùng lặp ảnh trong gói OpenXML (.xlsx) để tối ưu dung lượng."""
    if not xlsx_path or not os.path.isfile(xlsx_path) or not str(xlsx_path).lower().endswith('.xlsx'):
        return False
    try:
        with zipfile.ZipFile(xlsx_path, 'r') as zin:
            names = zin.namelist()
            media_files = [n for n in names if n.startswith('xl/media/')]
            if len(media_files) <= 1:
                return True

            by_hash = {}
            for mf in media_files:
                data = zin.read(mf)
                h = hashlib.sha256(data).hexdigest()
                by_hash.setdefault(h, []).append(mf)

            dup_map = {}
            canonical_files = set()
            for h, files in by_hash.items():
                files_sorted = sorted(files, key=lambda x: (len(os.path.splitext(os.path.basename(x))[0]), x))
                canon = files_sorted[0]
                canonical_files.add(canon)
                canon_name = os.path.basename(canon)
                for f in files_sorted[1:]:
                    dup_map[os.path.basename(f)] = canon_name

            if not dup_map:
                return True

            new_contents = {}
            for name in names:
                if name.endswith('.rels'):
                    content = zin.read(name).decode('utf-8')
                    modified = False
                    for dup_name, canon_name in dup_map.items():
                        if dup_name in content:
                            content = content.replace(dup_name, canon_name)
                            modified = True
                    if modified:
                        new_contents[name] = content.encode('utf-8')

            temp_out = xlsx_path + f".tmp_dedup_{os.getpid()}.xlsx"
            with zipfile.ZipFile(temp_out, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    if item.filename.startswith('xl/media/') and item.filename not in canonical_files:
                        continue
                    if item.filename in new_contents:
                        zout.writestr(item.filename, new_contents[item.filename])
                    else:
                        zout.writestr(item, zin.read(item.filename))

        with zipfile.ZipFile(temp_out, 'r') as zcheck:
            if zcheck.testzip() is not None:
                if os.path.exists(temp_out):
                    os.remove(temp_out)
                return False

        os.replace(temp_out, xlsx_path)
        if verbose:
            print(f"[DedupMedia] Đã khử trùng lặp {len(dup_map)} ảnh thừa trong {os.path.basename(xlsx_path)}.")
        return True
    except Exception as e:
        if verbose:
            print(f"[DedupMedia] Bỏ qua tối ưu ảnh: {e}")
        return False


def find_cell_text(ws, search_text):
    """Tìm cell chứa text (case-insensitive, partial) trong used range."""
    search_lower = search_text.lower()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1,
                            min_col=1, max_col=ws.max_column or 1):
        for cell in row:
            if cell.value is not None and search_lower in str(cell.value).lower():
                return cell
    return None


def main():
    parser = argparse.ArgumentParser(description="Excel month-template cleanup (openpyxl)")
    parser.add_argument('--file', required=True, help="Path to Excel workbook")
    args = parser.parse_args()

    print("Starting Create-Month cleanup (Python)...")
    print(f"File: {args.file}")

    # Mở workbook
    try:
        wb = load_workbook(args.file)
    except Exception as e:
        print(f"ERROR: Failed to open workbook: {e}", file=sys.stderr)
        sys.exit(1)

    try:
        # 1. Thu thập các sheet đơn hàng (pattern "N-..."), chọn sheet số cao nhất làm template
        order_sheet_names = []
        max_num = -1
        template_name = None

        for name in wb.sheetnames:
            m = SHEET_PATTERN.match(name)
            if m:
                num = int(m.group(1))
                if num < MAX_SHEET_NUM:
                    order_sheet_names.append(name)
                    if num > max_num:
                        max_num = num
                        template_name = name

        if template_name is None:
            raise ValueError("No numbered order sheet found to use as template.")

        print(f"Template sheet chosen: {template_name}")

        # 2. Xóa tất cả sheet đơn hàng khác (chỉ giữ template)
        for name in order_sheet_names:
            if name != template_name:
                print(f"Deleting order sheet: {name}")
                del wb[name]

        ws = wb[template_name]

        # 3. Dọn template: xóa khách hàng / ngày + dòng sản phẩm
        cust_label = find_cell_text(ws, "Kh\u00e1ch h\u00e0ng")
        if cust_label is None:
            cust_label = find_cell_text(ws, "Khach hang")
        if cust_label is not None:
            ws.cell(row=cust_label.row + 1, column=cust_label.column).value = ""

        date_label = find_cell_text(ws, "Ng\u00e0y")
        if date_label is None:
            date_label = find_cell_text(ws, "Ngay")
        if date_label is not None:
            ws.cell(row=date_label.row + 1, column=date_label.column).value = ""

        # Tìm header "Sản phẩm" / "Product name"
        header_cell = find_cell_text(ws, "S\u1ea3n ph\u1ea9m")
        if header_cell is None:
            header_cell = find_cell_text(ws, "Product name")

        if header_cell is not None:
            header_row = header_cell.row
            start_row = header_row + 1

            # Tìm footer "Tổng" / "Total"
            footer_row = 0
            for r in range(start_row, start_row + 201):
                val = ws.cell(row=r, column=2).value
                if val is not None:
                    val_str = str(val).strip().lower()
                    if "t\u1ed5ng" in val_str or "total" in val_str:
                        footer_row = r
                        break

            if footer_row > 0:
                print(f"Clearing item rows between {start_row} and {footer_row - 1}...")
                for r in range(start_row, footer_row):
                    for c in range(1, (ws.max_column or 15) + 1):
                        cell = ws.cell(row=r, column=c)
                        # Chỉ xóa constant, giữ formula
                        if cell.value is not None and not (isinstance(cell.value, str) and cell.value.startswith('=')):
                            cell.value = None
            else:
                print("Footer row (Tong/Total) not found; skipping item-row clear.")
        else:
            print("Header (San pham/Product name) not found; skipping item-row clear.")

        # 4. Đổi tên template thành "0-MAU"
        ws.title = "0-MAU"
        print("Template renamed to 0-MAU.")

        # Lưu workbook
        wb.save(args.file)
        print("Workbook saved successfully.")

        # 5. Heal: openpyxl output can require an Excel "repair" on open,
        #    which breaks the background automation. Re-save through Excel
        #    COM now so the fresh month file is clean from the start.
        heal_with_excel(args.file)

        # 6. Tự động khử trùng lặp ảnh media trong workbook
        deduplicate_media_images(args.file, verbose=True)

        print("CREATE_MONTH_OK")
        print("Create-Month cleanup (Python) completed.")

    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
